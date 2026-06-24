import io
import logging
from decimal import Decimal

from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, Query
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/truestat", tags=["unit_fact"])


COL_MAP = {
    0: "manager",
    1: "vendor_code",
    2: "size",
    3: "sales_count",
    4: "returns_count",
    5: "returns_pct",
    6: "net_sales",
    7: "revenue",
    8: "revenue_per_unit",
    9: "cost_price_total",
    10: "cost_price_per_unit",
    11: "cost_price_pct",
    12: "commission",
    13: "commission_pct",
    14: "logistics_total",
    15: "logistics_per_unit",
    16: "logistics_pct",
    17: "logistics_direct",
    18: "logistics_direct_pct",
    19: "logistics_return",
    20: "logistics_return_pct",
    21: "acquiring_penalty",
    22: "ad_spend",
    23: "ad_spend_per_unit",
    # col 24 (Y) is empty separator
    25: "margin_per_unit",
    26: "margin_pct",
    27: "roi",
    28: "orders_4w",
    29: "sales_4w",
    30: "buyout_4w_pct",
    31: "orders_18w",
    32: "sales_18w",
    33: "buyout_18w_pct",
    34: "stock_wb",
    35: "stock_in_transit",
    36: "stock_days",
    37: "stock_wb_prev",
    38: "avg_sales_per_week",
    39: "turnover_days",
}


def _to_float(v, default=0.0) -> float:
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if isinstance(v, Decimal):
        return round(float(v), 2)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", ".").replace("₽", "").replace("%", "")
    if not s or s == "-":
        return default
    try:
        return round(float(s), 2)
    except ValueError:
        return default


def _to_int(v) -> int:
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    if not s or s == "-":
        return 0
    try:
        return int(float(s))
    except ValueError:
        return 0


def _detect_row_type(a_val: str | None, b_val: str | None, c_val: str | None, row_idx: int) -> tuple[str, str, str, str]:
    """Returns (manager, vendor_code, size, row_type)."""
    a = str(a_val).strip() if a_val is not None else ""
    b = str(b_val).strip() if b_val is not None else ""
    c = str(c_val).strip() if c_val is not None else ""

    if a == "Итого" and not b:
        return ("", "", "", "grand_total")
    if a and "Всего" in a and not b:
        return (a.replace("Всего (", "").replace(")", "").strip(), "", "", "manager_total")
    if a and not b and not c:
        return (a, "", "", "manager_header")
    if b and b.startswith("Всего"):
        vc = b.replace("Всего (", "").replace(")", "").strip()
        return ("", vc, "", "vendor_total")
    if b and not c:
        return ("", b, "", "vendor_header")
    if b and c:
        return ("", b, c, "detail")
    return ("", "", "", "skip")


@router.post("/upload-unit-fact")
async def upload_unit_fact(file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    if not (file.filename or "").endswith(".xlsx"):
        raise HTTPException(400, "Only .xlsx files accepted")
    data = await file.read()
    if len(data) > 50 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 50MB)")

    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        raise HTTPException(400, "Empty workbook")

    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)

    rows_to_insert = []
    current_manager = ""

    for row_idx, row in enumerate(rows_iter, start=2):
        vals = list(row)
        if not vals or not any(v is not None for v in vals):
            continue

        a_val = vals[0] if len(vals) > 0 else None
        b_val = vals[1] if len(vals) > 1 else None
        c_val = vals[2] if len(vals) > 2 else None

        manager, vendor_code, size, row_type = _detect_row_type(a_val, b_val, c_val, row_idx)

        if row_type == "skip":
            continue

        if row_type == "manager_header":
            if manager:
                current_manager = manager
            continue

        if row_type == "manager_total" and manager:
            current_manager = manager
            continue

        if row_type == "vendor_header":
            rows_to_insert.append({
                "manager": current_manager,
                "vendor_code": vendor_code or "",
                "size": "",
                "row_type": "vendor_total",
                **{fld: 0.0 for fld in COL_MAP.values() if fld not in ("manager", "vendor_code", "size")},
            })
            continue

        if row_type == "vendor_total":
            b = b_val.strip() if b_val else ""
            vc = b.replace("Всего (", "").replace(")", "").strip()
            item = {"manager": current_manager, "vendor_code": vc, "size": "", "row_type": "vendor_total"}
            for col_idx, fld in COL_MAP.items():
                if fld in ("manager", "vendor_code", "size"):
                    continue
                v = vals[col_idx] if col_idx < len(vals) else None
                if fld in ("sales_count", "net_sales", "orders_4w", "sales_4w", "orders_18w", "sales_18w",
                           "stock_wb", "stock_in_transit", "stock_wb_prev"):
                    item[fld] = _to_int(v)
                else:
                    item[fld] = _to_float(v)
            rows_to_insert.append(item)
            continue

        item = {"manager": current_manager, "vendor_code": vendor_code or "", "size": size or "", "row_type": "detail"}
        for col_idx, fld in COL_MAP.items():
            if fld in ("manager", "vendor_code", "size"):
                continue
            v = vals[col_idx] if col_idx < len(vals) else None
            if fld in ("sales_count", "net_sales", "orders_4w", "sales_4w", "orders_18w", "sales_18w",
                       "stock_wb", "stock_in_transit", "stock_wb_prev"):
                item[fld] = _to_int(v)
            else:
                item[fld] = _to_float(v)
        rows_to_insert.append(item)

    if not rows_to_insert:
        raise HTTPException(400, "No data rows found")

    insert_cols = ["row_type"] + list(COL_MAP.values())
    placeholders = ", ".join(f":{f}" for f in insert_cols)
    columns = ", ".join(insert_cols)

    await db.execute(text("DELETE FROM unit_fact"))

    for item in rows_to_insert:
        params = {f: item.get(f, 0) for f in insert_cols}
        await db.execute(
            text(f"INSERT INTO unit_fact ({columns}) VALUES ({placeholders})"),
            params,
        )

    await db.commit()

    return {
        "status": "ok",
        "file": file.filename,
        "imported": len(rows_to_insert),
    }


@router.get("/unit-fact")
async def get_unit_fact(
    vendor_code: str = Query(default=""),
    row_type: str = Query(default=""),
    db: AsyncSession = Depends(get_db),
):
    where_clauses = []
    params = {}

    if vendor_code:
        where_clauses.append("vendor_code = :vendor_code")
        params["vendor_code"] = vendor_code
    if row_type:
        where_clauses.append("row_type = :row_type")
        params["row_type"] = row_type

    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)

    all_cols = ["id", "row_type", "manager", "vendor_code", "size",
                "sales_count", "returns_count", "returns_pct", "net_sales",
                "revenue", "revenue_per_unit", "cost_price_total", "cost_price_per_unit",
                "cost_price_pct", "commission", "commission_pct",
                "logistics_total", "logistics_per_unit", "logistics_pct",
                "logistics_direct", "logistics_direct_pct",
                "logistics_return", "logistics_return_pct",
                "acquiring_penalty", "ad_spend", "ad_spend_per_unit",
                "margin_per_unit", "margin_pct", "roi",
                "orders_4w", "sales_4w", "buyout_4w_pct",
                "orders_18w", "sales_18w", "buyout_18w_pct",
                "stock_wb", "stock_in_transit", "stock_days",
                "stock_wb_prev", "avg_sales_per_week", "turnover_days",
                "created_at", "updated_at"]
    cols_sql = ", ".join(all_cols)

    sql = text(f"SELECT {cols_sql} FROM unit_fact {where_sql} ORDER BY id ASC")
    result = await db.execute(sql, params)
    rows = result.fetchall()

    data = []
    for row in rows:
        data.append(dict(zip(all_cols, row)))

    return {"data": data, "count": len(data)}
